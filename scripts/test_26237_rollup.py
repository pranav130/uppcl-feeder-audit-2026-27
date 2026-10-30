import openpyxl, re
from collections import defaultdict

SIR_FILE = r'G:\Other computers\Prabhakar Sir\G Drive\FEEDERWISE ENERGY AUDIT\new\FY2026-27\11KV total feeders prog audit 2026-27 upto sep-26 F_excluding absurd CA.xlsx'
ENERGY_XLSX = r'C:\Users\HP\Documents\New Sep Input energy data.xlsx'

def norm(s):
    return re.sub(r'[^A-Z0-9]', '', str(s or '').strip().upper())

# Load 26237 list
wb_e = openpyxl.load_workbook(ENERGY_XLSX, read_only=True, data_only=True)
ws_e = wb_e['Sheet1']
feeders_26237 = []
for i, row in enumerate(ws_e.iter_rows(values_only=True)):
    if i == 0: continue
    if row[9] is True or str(row[9]).upper() == 'TRUE':
        feeders_26237.append(row)
wb_e.close()

# Load Sir sheet
wb_s = openpyxl.load_workbook(SIR_FILE, read_only=True, data_only=True)
ws_s = wb_s['AUDIT DATA']
sir_by_div_ss_fdr = {}
sir_by_ss_fdr = {}
sir_by_fdr = {}
for i, row in enumerate(ws_s.iter_rows(values_only=True)):
    if i < 2: continue
    if row[0] is None and row[6] is None: continue
    div = norm(row[4]); ss = norm(row[5]); fdr = norm(row[6])
    sir_by_div_ss_fdr[div + ss + fdr] = row
    sir_by_ss_fdr[ss + fdr] = row
    if fdr not in sir_by_fdr: sir_by_fdr[fdr] = row
wb_s.close()

discom_totals = defaultdict(lambda: {'tot': 0, 'll': defaultdict(int), 'atc': defaultdict(int)})
uppcl_total = {'tot': 0, 'll': defaultdict(int), 'atc': defaultdict(int)}

for r in feeders_26237:
    div = norm(r[3]); ss = norm(r[4]); fdr = norm(r[8])
    sr = sir_by_div_ss_fdr.get(div + ss + fdr) or sir_by_ss_fdr.get(ss + fdr) or sir_by_fdr.get(fdr)
    
    d = str(sr[0]).strip().upper()
    ll_slab = str(sr[18]).strip()
    atc_slab = str(sr[19]).strip()
    
    discom_totals[d]['tot'] += 1
    discom_totals[d]['ll'][ll_slab] += 1
    discom_totals[d]['atc'][atc_slab] += 1
    
    uppcl_total['tot'] += 1
    uppcl_total['ll'][ll_slab] += 1
    uppcl_total['atc'][atc_slab] += 1

print(f"Total 26237 UPPCL count: {uppcl_total['tot']}")
for d, info in sorted(discom_totals.items()):
    print(f"{d}: {info['tot']} feeders")

print("\nUPPCL Total Line Loss Slabs (on 26,237):")
for s, c in sorted(uppcl_total['ll'].items()):
    print(f"  {s}: {c}")

print("\nUPPCL Total AT&C Loss Slabs (on 26,237):")
for s, c in sorted(uppcl_total['atc'].items()):
    print(f"  {s}: {c}")
