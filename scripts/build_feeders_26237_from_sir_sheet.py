import openpyxl
import json
import os
import re

SIR_FILE = r'G:\Other computers\Prabhakar Sir\G Drive\FEEDERWISE ENERGY AUDIT\new\FY2026-27\11KV total feeders prog audit 2026-27 upto sep-26 F_excluding absurd CA.xlsx'
ENERGY_XLSX = r'C:\Users\HP\Documents\New Sep Input energy data.xlsx'
REPO_DIR = r"C:\Users\HP\.gemini\antigravity\scratch\UPPCL_11KV_PSR_Dashboard_2026_27"

def clean_num(val):
    if val is None or val == '' or str(val).strip() in ('', 'None', '-'): return 0.0
    try: return round(float(val), 2)
    except: return 0.0

def clean_num1(val):
    if val is None or val == '' or str(val).strip() in ('', 'None', '-'): return 0.0
    try: return round(float(val), 1)
    except: return 0.0

def clean_int(val):
    if val is None or val == '' or str(val).strip() in ('', 'None', '-'): return 0
    try: return int(round(float(val)))
    except: return 0

def clean_str(val):
    if val is None: return ''
    return str(val).strip()

def norm(s):
    return re.sub(r'[^A-Z0-9]', '', str(s or '').strip().upper())

# 1. Load 26,237 feeders in original order
print("Loading 26,237 feeders from New Sep Input energy data.xlsx...")
wb_e = openpyxl.load_workbook(ENERGY_XLSX, read_only=True, data_only=True)
ws_e = wb_e['Sheet1']
feeders_26237 = []
for i, row in enumerate(ws_e.iter_rows(values_only=True)):
    if i == 0: continue
    if row[9] is True or str(row[9]).upper() == 'TRUE':
        feeders_26237.append(row)
wb_e.close()
print(f"Loaded {len(feeders_26237)} valid feeders.")

# 2. Load Sir AUDIT DATA
print("Loading Sir sheet AUDIT DATA...")
wb_s = openpyxl.load_workbook(SIR_FILE, read_only=True, data_only=True)
ws_s = wb_s['AUDIT DATA']

sir_by_div_ss_fdr = {}
sir_by_ss_fdr = {}
sir_by_fdr = {}

for i, row in enumerate(ws_s.iter_rows(values_only=True)):
    if i < 2: continue
    if row[0] is None and row[6] is None: continue
    
    div = norm(row[4])
    ss = norm(row[5])
    fdr = norm(row[6])
    
    sir_by_div_ss_fdr[div + ss + fdr] = row
    sir_by_ss_fdr[ss + fdr] = row
    if fdr not in sir_by_fdr:
        sir_by_fdr[fdr] = row

wb_s.close()
print("Sir AUDIT DATA indexed.")

# Dimension lists
discoms = ['PASCHIMANCHAL', 'DAKSHINANCHAL', 'POORVANCHAL', 'MADHYANCHAL', 'KESCO']
areas = ['URBAN', 'TEHSIL', 'RURAL']
natures = [
    'OTHER', 'INDEPENDENT', 'MIXED ABOVE 50%', 'SPARE',
    'SUBSTATION', 'MIXED ABOVE 25%', 'AGRICULTURE', 'INDUSTRIAL ABOVE 75%'
]

d_map = {d: i for i, d in enumerate(discoms)}
a_map = {a: i for i, a in enumerate(areas)}
n_map = {n: i for i, n in enumerate(natures)}

zones = []; z_map = {}
circles = []; c_map = {}
divisions = []; div_map = {}
substations = []; ss_map = {}

def get_or_add(val, lst, mp):
    s = clean_str(val).upper() or 'UNKNOWN'
    if s not in mp:
        mp[s] = len(lst)
        lst.append(s)
    return mp[s]

def normalize_area(a):
    u = clean_str(a).upper()
    if 'TEHSIL' in u: return 'TEHSIL'
    if 'RURAL' in u: return 'RURAL'
    return 'URBAN'

def normalize_nature(n):
    u = clean_str(n).upper()
    for valid in natures:
        if valid == u: return valid
    for valid in natures:
        if valid in u: return valid
    return 'OTHER'

output_rows = []

for r in feeders_26237:
    div = norm(r[3])
    ss = norm(r[4])
    fdr = norm(r[8])
    
    sr = sir_by_div_ss_fdr.get(div + ss + fdr) or sir_by_ss_fdr.get(ss + fdr) or sir_by_fdr.get(fdr)
    if not sr:
        print(f"ERROR: Feeder {r[8]} not found in Sir sheet!")
        continue
        
    disc = clean_str(sr[0]).upper()
    zone = clean_str(sr[2]).upper()
    circle = clean_str(sr[3]).upper()
    division = clean_str(sr[4]).upper()
    substation = clean_str(sr[5]).upper()
    feeder = clean_str(sr[6])
    
    area = normalize_area(sr[8])
    nature = normalize_nature(sr[9])
    
    d_idx = d_map.get(disc, 0)
    z_idx = get_or_add(zone, zones, z_map)
    c_idx = get_or_add(circle, circles, c_map)
    div_idx = get_or_add(division, divisions, div_map)
    ss_idx = get_or_add(substation, substations, ss_map)
    a_idx = a_map[area]
    n_idx = n_map[nature]
    
    # Monthly IE (Cols 28-33: Apr-Sep) - 2 decimal places to preserve fractional metering (e.g. 0.01, 0.02 MU)
    ie = [clean_num(sr[c]) for c in range(28, 34)]
    
    # Monthly SE (Cols 52-57: Apr-Sep)
    se = [clean_num(sr[c]) for c in range(52, 58)]
    
    # Monthly Ass (Cols 64-69: Apr-Sep)
    ass = [clean_num(sr[c]) for c in range(64, 70)]
    
    # Monthly Real (Cols 76-81: Apr-Sep)
    real = [clean_num(sr[c]) for c in range(76, 82)]
    
    # Monthly PTW Ass (Cols 92-97)
    ptw_ass = [clean_num(sr[c]) for c in range(92, 98)]
    
    # Monthly PTW Subs (Cols 104-109)
    ptw_sub = [clean_num(sr[c]) for c in range(104, 110)]
    
    # Billable Sep'26 (Col 45)
    cons = clean_int(sr[45])
    
    tot_ie = sum(ie)
    tot_se = sum(se)
    
    flags = 0
    if tot_ie > 0:
        ll_pct = ((tot_ie - tot_se) / tot_ie) * 100.0
        if ll_pct < 0 or ll_pct > 100:
            flags |= 1
        if tot_se > tot_ie:
            flags |= 2
    else:
        flags |= 4
        
    if cons == 0:
        flags |= 8
        
    row_compact = [
        d_idx, z_idx, c_idx, div_idx, ss_idx, feeder,
        a_idx, n_idx, ie, se, ass, real,
        cons, flags, ptw_ass, ptw_sub
    ]
    output_rows.append(row_compact)

assert len(output_rows) == 26237, f"Expected 26237 rows, got {len(output_rows)}"
print(f"SUCCESS: Extracted exactly {len(output_rows)} rows!")

dataset = {
    'd': {
        'Discom': discoms,
        'Zone': zones,
        'Circle': circles,
        'Division': divisions,
        'Area': areas,
        'Nature': natures,
        'SS': substations
    },
    'rows': output_rows
}

for folder in ['data', 'frontend/data']:
    target = os.path.join(REPO_DIR, folder, 'feeders.json')
    with open(target, 'w', encoding='utf-8') as f:
        json.dump(dataset, f, separators=(',', ':'))
    print(f"Wrote {target} ({os.path.getsize(target)/1024/1024:.2f} MB)")
